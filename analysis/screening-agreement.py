#!/usr/bin/env python3
"""Agreement between the first author's re-screening decisions and the
second coder's blind decisions on the 12-paper sample.

Reports raw agreement and Cohen's kappa on Include/Exclude, and agreement
on the exclusion code where both coders excluded.

Usage:
    python3 analysis/screening-agreement.py \
        docs/screening/rescreening-decisions.xlsx \
        docs/screening/second-coder-sample.xlsx
Requires openpyxl.
"""

import sys
from collections import Counter

import openpyxl


def rows(path, sheet):
    ws = openpyxl.load_workbook(path, data_only=True)[sheet]
    header = [c.value for c in ws[1]]
    return [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]


def norm_decision(v):
    v = (v or "").strip().lower()
    if v.startswith("include"):
        return "Include", ""
    if v.startswith("exclude"):
        code = v.split("(")[-1].rstrip(")").upper() if "(" in v else ""
        return "Exclude", code
    return None, ""


def main():
    first_path, second_path = sys.argv[1], sys.argv[2]
    first = {}
    for r in rows(first_path, "Rescreen_58"):
        first[r["Paper ID"]] = norm_decision(r["AUTHOR DECISION"])
    pairs = []
    for r in rows(second_path, "Second_Coder"):
        d2, _ = norm_decision(r["Decision (Include/Exclude)"])
        c2 = (r["Exclusion code (E1/E2/E3/SCOPE, blank if included)"] or "").strip().upper()
        d1, c1 = first[r["Paper ID"]]
        if d1 is None or d2 is None:
            sys.exit(f"Missing decision for {r['Paper ID']}")
        pairs.append((r["Paper ID"], d1, c1, d2, c2))

    n = len(pairs)
    agree = sum(1 for _, d1, _, d2, _ in pairs if d1 == d2)
    p_o = agree / n
    a = Counter(d1 for _, d1, _, _, _ in pairs)
    b = Counter(d2 for _, _, _, d2, _ in pairs)
    p_e = sum(a[k] * b[k] for k in ("Include", "Exclude")) / n ** 2
    kappa = (p_o - p_e) / (1 - p_e) if p_e < 1 else float("nan")

    both_ex = [(c1, c2) for _, d1, c1, d2, c2 in pairs if d1 == d2 == "Exclude"]
    code_agree = sum(1 for c1, c2 in both_ex if c1 == c2)

    print(f"Papers: {n}")
    print(f"Decision agreement: {agree}/{n} ({100 * p_o:.1f}%)")
    print(f"Cohen's kappa (Include/Exclude): {kappa:.2f}")
    print(f"Exclusion-code agreement where both excluded: {code_agree}/{len(both_ex)}")
    for pid, d1, c1, d2, c2 in pairs:
        if d1 != d2 or (d1 == "Exclude" and c1 != c2):
            print(f"  Disagreement {pid}: first {d1} {c1}, second {d2} {c2}")


if __name__ == "__main__":
    main()
